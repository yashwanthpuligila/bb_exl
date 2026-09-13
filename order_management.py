import argparse
import os
import time
from datetime import datetime, date
from typing import List, Dict, Any, Iterable, Optional, Tuple

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install with 'pip install -r requirements.txt'"
    ) from exc


WORKBOOK_FILENAME: str = "orders.xlsx"
SHEET_NAME: str = "Orders"
HEADERS: List[str] = [
    "order_id",
    "timestamp",
    "customer_name",
    "contact",
    "product_name",
    "quantity",
    "price",
    "total_price",
]


def is_file_accessible(filepath: str, max_attempts: int = 3) -> bool:
    """Check if file is accessible for writing, with retry logic."""
    for attempt in range(max_attempts):
        try:
            # Try to open file in append mode to check write access
            with open(filepath, 'a'):
                pass
            return True
        except (PermissionError, IOError):
            if attempt < max_attempts - 1:
                time.sleep(0.5)  # Wait a bit before retrying
            continue
    return False


def check_excel_lock() -> Optional[str]:
    """Check if Excel has a lock on the orders file."""
    lock_file = f"~${WORKBOOK_FILENAME}"
    if os.path.exists(lock_file):
        return ("Excel lock file detected. Please close Excel if you have orders.xlsx open, "
                "then try again.")
    return None


def ensure_workbook_exists() -> None:
    if not os.path.exists(WORKBOOK_FILENAME):
        # Check if we can create the file
        if not is_file_accessible(WORKBOOK_FILENAME):
            raise PermissionError(f"Cannot create {WORKBOOK_FILENAME}. Check file permissions.")
        
        wb: Workbook = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(HEADERS)
        wb.save(WORKBOOK_FILENAME)
    else:
        # Check for Excel lock
        lock_msg = check_excel_lock()
        if lock_msg:
            raise PermissionError(lock_msg)
        
        # Check if existing file is accessible
        if not is_file_accessible(WORKBOOK_FILENAME):
            raise PermissionError(f"Cannot access {WORKBOOK_FILENAME}. File may be open in Excel or you may lack permissions.")


def open_sheet():
    ensure_workbook_exists()
    
    try:
        wb = load_workbook(WORKBOOK_FILENAME)
    except PermissionError:
        lock_msg = check_excel_lock()
        if lock_msg:
            raise PermissionError(lock_msg)
        raise PermissionError(f"Cannot open {WORKBOOK_FILENAME}. File may be open in another program.")
    
    if SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_NAME)
        ws.append(HEADERS)
        wb.save(WORKBOOK_FILENAME)
    ws = wb[SHEET_NAME]
    return wb, ws


def next_order_id(current_max_row: int) -> int:
    # current_max_row includes the header row.
    return max(1, current_max_row - 1 + 1)


def append_order(
    customer_name: str,
    contact: str,
    product_name: str,
    quantity: int,
    price: float,
) -> Tuple[int, float]:
    wb, ws = open_sheet()
    oid = next_order_id(ws.max_row)
    ts = datetime.now().isoformat(timespec="seconds")
    total = float(quantity) * float(price)
    ws.append([oid, ts, customer_name, contact, product_name, quantity, price, total])
    wb.save(WORKBOOK_FILENAME)
    return oid, total


def iter_orders() -> Iterable[Dict[str, Any]]:
    _, ws = open_sheet()
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            continue  # skip header
        if not any(cell is not None for cell in row):
            continue
        yield {
            "order_id": row[0],
            "timestamp": row[1],
            "customer_name": row[2] or "",
            "contact": row[3] or "",
            "product_name": row[4] or "",
            "quantity": int(row[5]) if row[5] is not None else 0,
            "price": float(row[6]) if row[6] is not None else 0.0,
            "total_price": float(row[7]) if row[7] is not None else 0.0,
        }


def print_table(rows: List[Dict[str, Any]], columns: List[str]) -> None:
    if not rows:
        print("No records found.")
        return
    # Compute column widths
    widths = {col: len(col) for col in columns}
    for row in rows:
        for col in columns:
            val = row.get(col, "")
            text = f"{val}"
            if len(text) > widths[col]:
                widths[col] = len(text)

    def fmt_row(cells: List[str]) -> str:
        return " | ".join(
            str(cell).ljust(widths[columns[idx]]) for idx, cell in enumerate(cells)
        )

    # Header
    header_line = fmt_row(columns)
    sep = "-+-".join("-" * widths[c] for c in columns)
    print(header_line)
    print(sep)
    # Rows
    for row in rows:
        print(fmt_row([row.get(c, "") for c in columns]))


def filter_orders(
    orders: Iterable[Dict[str, Any]],
    customer_query: Optional[str] = None,
    product_query: Optional[str] = None,
    since: Optional[date] = None,
    until: Optional[date] = None,
) -> List[Dict[str, Any]]:
    def contains(value: str, query: str) -> bool:
        return query.lower() in (value or "").lower()

    filtered: List[Dict[str, Any]] = []
    for order in orders:
        if customer_query and not contains(order.get("customer_name", ""), customer_query):
            continue
        if product_query and not contains(order.get("product_name", ""), product_query):
            continue
        if since or until:
            try:
                ts = order.get("timestamp", "")
                dt = datetime.fromisoformat(str(ts)).date()
            except Exception:
                continue
            if since and dt < since:
                continue
            if until and dt > until:
                continue
        filtered.append(order)
    return filtered


def summarize_sales(orders: Iterable[Dict[str, Any]], period: str) -> List[Dict[str, Any]]:
    buckets: Dict[str, float] = {}
    for order in orders:
        try:
            ts = order.get("timestamp", "")
            dt = datetime.fromisoformat(str(ts))
        except Exception:
            continue
        if period == "daily":
            key = dt.date().isoformat()
        elif period == "weekly":
            iso = dt.isocalendar()  # (year, week, weekday)
            key = f"{iso.year}-W{iso.week:02d}"
        elif period == "monthly":
            key = f"{dt.year}-{dt.month:02d}"
        else:
            raise ValueError("Invalid period")
        buckets[key] = buckets.get(key, 0.0) + float(order.get("total_price", 0.0))
    # Convert to rows
    rows = [
        {"period": k, "sales_total": round(v, 2)} for k, v in sorted(buckets.items())
    ]
    return rows


def export_orders(rows: List[Dict[str, Any]], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for row in rows:
        ws.append(
            [
                row.get("order_id"),
                row.get("timestamp"),
                row.get("customer_name"),
                row.get("contact"),
                row.get("product_name"),
                row.get("quantity"),
                row.get("price"),
                row.get("total_price"),
            ]
        )
    wb.save(output_path)


def parse_date_str(s: Optional[str]) -> Optional[date]:
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        raise SystemExit("Invalid date format. Use YYYY-MM-DD.")


def cmd_add(args: argparse.Namespace) -> None:
    oid, total = append_order(
        customer_name=args.customer,
        contact=args.contact or "",
        product_name=args.product,
        quantity=int(args.quantity),
        price=float(args.price),
    )
    print(
        f"Saved order #{oid} for {args.customer} | {args.product} x{args.quantity} => {total:.2f}"
    )


def cmd_list(_: argparse.Namespace) -> None:
    rows = list(iter_orders())
    print_table(rows, [
        "order_id",
        "timestamp",
        "customer_name",
        "contact",
        "product_name",
        "quantity",
        "price",
        "total_price",
    ])
    total_sales = sum(r.get("total_price", 0.0) for r in rows)
    print(f"\nOrders: {len(rows)} | Sales total: {total_sales:.2f}")


def cmd_search(args: argparse.Namespace) -> None:
    if not args.customer and not args.product:
        raise SystemExit("Provide --customer and/or --product for searching.")
    rows = filter_orders(
        iter_orders(),
        customer_query=args.customer,
        product_query=args.product,
    )
    print_table(rows, [
        "order_id",
        "timestamp",
        "customer_name",
        "contact",
        "product_name",
        "quantity",
        "price",
        "total_price",
    ])
    total_sales = sum(r.get("total_price", 0.0) for r in rows)
    print(f"\nMatched: {len(rows)} | Sales total: {total_sales:.2f}")


def cmd_dashboard(args: argparse.Namespace) -> None:
    rows = list(iter_orders())
    summary = summarize_sales(rows, args.period)
    print_table(summary, ["period", "sales_total"])


def cmd_export(args: argparse.Namespace) -> None:
    since = parse_date_str(args.since)
    until = parse_date_str(args.until)
    rows = filter_orders(
        iter_orders(),
        customer_query=args.customer,
        product_query=args.product,
        since=since,
        until=until,
    )
    output = args.output or "export.xlsx"
    export_orders(rows, output)
    print(f"Exported {len(rows)} rows to {output}")


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description=(
            "Order Management CLI (Excel-backed). Stores customer & order details in a single sheet."
        )
    )
    sub = p.add_subparsers(dest="command", required=True)

    add = sub.add_parser("add", help="Add a new order")
    add.add_argument("--customer", required=True, help="Customer name")
    add.add_argument("--contact", required=False, help="Contact info (phone/email)")
    add.add_argument("--product", required=True, help="Product name")
    add.add_argument("--quantity", required=True, type=int, help="Quantity")
    add.add_argument("--price", required=True, type=float, help="Unit price")
    add.set_defaults(func=cmd_add)

    lst = sub.add_parser("list", help="List all past orders")
    lst.set_defaults(func=cmd_list)

    sch = sub.add_parser("search", help="Search by customer or product name")
    sch.add_argument("--customer", required=False, help="Customer name contains")
    sch.add_argument("--product", required=False, help="Product name contains")
    sch.set_defaults(func=cmd_search)

    dash = sub.add_parser("dashboard", help="Sales summary by period")
    dash.add_argument(
        "--period",
        choices=["daily", "weekly", "monthly"],
        default="daily",
        help="Aggregation period",
    )
    dash.set_defaults(func=cmd_dashboard)

    exp = sub.add_parser("export", help="Export data to a new Excel file")
    exp.add_argument("--customer", required=False, help="Filter: customer contains")
    exp.add_argument("--product", required=False, help="Filter: product contains")
    exp.add_argument("--since", required=False, help="Filter: since date YYYY-MM-DD")
    exp.add_argument("--until", required=False, help="Filter: until date YYYY-MM-DD")
    exp.add_argument("-o", "--output", required=False, help="Output .xlsx path (default export.xlsx)")
    exp.set_defaults(func=cmd_export)

    return p


def main(argv: Optional[List[str]] = None) -> None:
    parser = build_parser()
    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()


