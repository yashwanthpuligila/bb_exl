import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import glob


class SimpleInvoiceApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Simple Invoice Generator")
        self.root.geometry("600x500")
        
        self.customer_info = {}
        self.products = []
        
        self.create_interface()
        
    def create_interface(self):
        main_frame = ttk.Frame(self.root, padding=20)
        main_frame.pack(fill='both', expand=True)
        
        # Title
        ttk.Label(main_frame, text="Simple Invoice Generator", 
                 font=('Arial', 16, 'bold')).pack(pady=(0, 20))
        
        # Customer Section
        customer_frame = ttk.LabelFrame(main_frame, text="Customer Information", padding=10)
        customer_frame.pack(fill='x', pady=(0, 15))
        
        ttk.Label(customer_frame, text="Shop Name:").grid(row=0, column=0, sticky='w', pady=3)
        self.shop_entry = ttk.Entry(customer_frame, width=30)
        self.shop_entry.grid(row=0, column=1, padx=(10, 0), pady=3)
        
        ttk.Label(customer_frame, text="Customer Name:").grid(row=1, column=0, sticky='w', pady=3)
        self.customer_entry = ttk.Entry(customer_frame, width=30)
        self.customer_entry.grid(row=1, column=1, padx=(10, 0), pady=3)
        
        ttk.Label(customer_frame, text="Area:").grid(row=2, column=0, sticky='w', pady=3)
        self.area_entry = ttk.Entry(customer_frame, width=30)
        self.area_entry.grid(row=2, column=1, padx=(10, 0), pady=3)
        
        # Product Section
        product_frame = ttk.LabelFrame(main_frame, text="Add Product", padding=10)
        product_frame.pack(fill='x', pady=(0, 15))
        
        ttk.Label(product_frame, text="Product:").grid(row=0, column=0, sticky='w', pady=3)
        self.product_entry = ttk.Entry(product_frame, width=25)
        self.product_entry.grid(row=0, column=1, padx=(5, 0), pady=3)
        
        ttk.Label(product_frame, text="Qty:").grid(row=0, column=2, sticky='w', pady=3, padx=(10, 0))
        self.qty_entry = ttk.Entry(product_frame, width=10)
        self.qty_entry.grid(row=0, column=3, padx=(5, 0), pady=3)
        
        ttk.Label(product_frame, text="Price:").grid(row=1, column=0, sticky='w', pady=3)
        self.price_entry = ttk.Entry(product_frame, width=15)
        self.price_entry.grid(row=1, column=1, padx=(5, 0), pady=3)
        
        ttk.Button(product_frame, text="Add Product", 
                  command=self.add_product).grid(row=1, column=2, columnspan=2, pady=5, padx=(10, 0))
        
        # Products List
        list_frame = ttk.LabelFrame(main_frame, text="Products Added", padding=10)
        list_frame.pack(fill='both', expand=True, pady=(0, 15))
        
        self.product_listbox = tk.Listbox(list_frame, height=6)
        self.product_listbox.pack(fill='both', expand=True)
        
        # Total and Generate
        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill='x')
        
        self.total_label = ttk.Label(bottom_frame, text="Total: $0.00", 
                                    font=('Arial', 12, 'bold'))
        self.total_label.pack(side='left')
        
        # Buttons frame
        buttons_frame = ttk.Frame(bottom_frame)
        buttons_frame.pack(side='right')
        
        ttk.Button(buttons_frame, text="New Customer", 
                  command=self.reset_form).pack(side='left', padx=(0, 10))
        
        ttk.Button(buttons_frame, text="Generate Invoice", 
                  command=self.generate_invoice,
                  style='Accent.TButton').pack(side='left')
        
        # Status
        self.status_label = ttk.Label(main_frame, text="Ready...", foreground='green')
        self.status_label.pack(pady=(10, 0))
        
    def add_product(self):
        try:
            product = self.product_entry.get().strip()
            qty = int(self.qty_entry.get().strip())
            price = float(self.price_entry.get().strip())
            
            if not product or qty <= 0 or price < 0:
                raise ValueError("Invalid input")
            
            total = qty * price
            
            product_info = {
                'name': product,
                'qty': qty,
                'price': price,
                'total': total
            }
            
            self.products.append(product_info)
            
            # Add to listbox
            display_text = f"{product} | Qty: {qty} | ${price:.2f} | Total: ${total:.2f}"
            self.product_listbox.insert(tk.END, display_text)
            
            # Clear entries
            self.product_entry.delete(0, tk.END)
            self.qty_entry.delete(0, tk.END)
            self.price_entry.delete(0, tk.END)
            
            # Update total
            self.update_total()
            
            self.status_label.config(text=f"Added: {product}", foreground='green')
            
        except ValueError:
            messagebox.showerror("Error", "Please enter valid product details")
            
    def update_total(self):
        total = sum(p['total'] for p in self.products)
        self.total_label.config(text=f"Total: ${total:.2f}")
        
    def generate_invoice(self):
        try:
            # Get customer info
            shop = self.shop_entry.get().strip()
            customer = self.customer_entry.get().strip()
            area = self.area_entry.get().strip()
            
            if not shop or not customer or not area:
                messagebox.showerror("Error", "Please fill customer information")
                return
                
            if not self.products:
                messagebox.showerror("Error", "Please add at least one product")
                return
            
            print(f"DEBUG: Starting invoice generation...")
            print(f"DEBUG: Customer: {customer}, Shop: {shop}, Area: {area}")
            print(f"DEBUG: Products: {len(self.products)}")
            
            # Create safe filename
            safe_name = "".join(c for c in customer if c.isalnum() or c.isspace()).strip().replace(' ', '_')
            customer_pattern = f"Invoice_{safe_name}_*.xlsx"
            
            # Check if customer file already exists
            existing_files = glob.glob(customer_pattern)
            current_date = datetime.now().strftime('%d/%m/%Y')
            current_time = datetime.now().strftime('%H:%M')
            
            if existing_files:
                # Use existing file (get the most recent one)
                existing_file = max(existing_files, key=os.path.getctime)
                print(f"DEBUG: Found existing file for {customer}: {existing_file}")
                
                # Get file creation date for reference
                file_creation_date = datetime.fromtimestamp(os.path.getctime(existing_file)).strftime('%d/%m/%Y')
                
                # Load existing workbook
                wb = load_workbook(existing_file)
                ws = wb.active
                
                # Find the last row with data
                last_row = ws.max_row
                
                # Find where to start adding new items (after previous total)
                new_start_row = last_row + 4  # More space for better separation
                
                # Add prominent separator for new order
                separator_row = new_start_row - 2
                ws.cell(row=separator_row, column=1, value=f"╔══════════════════════════════════════════════════════════════════════════════╗")
                ws.cell(row=separator_row, column=1).font = Font(name='Courier New', size=10, bold=True, color='0066CC')
                
                ws.cell(row=separator_row + 1, column=1, value=f"║ 📅 NEW ORDER - {current_date} at {current_time} (Previous orders from {file_creation_date}) ║")
                ws.cell(row=separator_row + 1, column=1).font = Font(name='Arial', size=11, bold=True, color='0066CC')
                
                ws.cell(row=separator_row + 2, column=1, value=f"╚══════════════════════════════════════════════════════════════════════════════╝")
                ws.cell(row=separator_row + 2, column=1).font = Font(name='Courier New', size=10, bold=True, color='0066CC')
                
                filename = existing_file
                
            else:
                # Create new workbook
                print(f"DEBUG: Creating new file for {customer}")
                wb = Workbook()
                ws = wb.active
                ws.title = "Invoice"
                
                # Company info
                ws['A1'] = "Your Company"
                ws['A1'].font = Font(size=16, bold=True, color='4472C4')
                ws['A2'] = "123 Your Address"
                ws['A3'] = "City, State 12345"
                ws['A4'] = "(555) 123-4567"
                
                # Invoice title with date
                ws['A6'] = f"Invoice for {shop}"
                ws['A6'].font = Font(size=20, bold=True, color='4472C4')
                
                ws['A7'] = f"📅 Started: {current_date} at {current_time}"
                ws['A7'].font = Font(name='Arial', size=10, color='FF0000')
                
                # Customer info
                ws['A9'] = "Customer Details:"
                ws['A9'].font = Font(bold=True)
                ws['A10'] = f"Shop: {shop}"
                ws['A11'] = f"Owner: {customer}"
                ws['A12'] = f"Area: {area}"
                
                new_start_row = 14
                filename = f"Invoice_{safe_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            print(f"DEBUG: Using file: {filename}")
            
            # Styling
            header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            normal_font = Font(name='Arial', size=10)
            total_font = Font(name='Arial', size=12, bold=True, color='FF0000')
            grand_total_font = Font(name='Arial', size=16, bold=True, color='FFFFFF')
            
            # Borders
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            
            thick_border = Border(
                left=Side(style='thick'), right=Side(style='thick'),
                top=Side(style='thick'), bottom=Side(style='thick')
            )
            
            # Headers with blue background
            headers = ["Product", "Qty", "Price", "Total"]
            for i, header in enumerate(headers, 1):
                cell = ws.cell(row=new_start_row, column=i)
                cell.value = header
                cell.font = header_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            
            # Products
            row = new_start_row + 1
            order_total = 0
            for product in self.products:
                ws.cell(row=row, column=1, value=product['name'])
                ws.cell(row=row, column=2, value=product['qty']).alignment = Alignment(horizontal='center')
                ws.cell(row=row, column=3, value=product['price'])
                ws.cell(row=row, column=4, value=product['total'])
                
                # Add borders
                for col in range(1, 5):
                    ws.cell(row=row, column=col).border = thin_border
                    ws.cell(row=row, column=col).font = normal_font
                
                order_total += product['total']
                row += 1
            
            # Order Total with date/time
            total_row = row + 1
            ws.cell(row=total_row, column=2, value=f"Order Total ({current_date} {current_time}):")
            ws.cell(row=total_row, column=2).font = Font(name='Arial', size=10, bold=True)
            ws.cell(row=total_row, column=2).alignment = Alignment(horizontal='right')
            
            ws.cell(row=total_row, column=4, value=order_total)
            ws.cell(row=total_row, column=4).font = total_font
            ws.cell(row=total_row, column=4).alignment = Alignment(horizontal='right')
            ws.cell(row=total_row, column=4).border = thick_border
            
            # Calculate grand total (sum all order totals in the file)
            grand_total = 0
            order_count = 0
            for row_num in range(1, ws.max_row + 1):
                cell_value = ws.cell(row=row_num, column=2).value
                if isinstance(cell_value, str) and "Order Total (" in cell_value:
                    total_cell = ws.cell(row=row_num, column=4)
                    if isinstance(total_cell.value, (int, float)):
                        grand_total += total_cell.value
                        order_count += 1
            
            # Grand Total (always show, even for single orders)
            grand_total_row = total_row + 3
            
            # Add visual separator before grand total
            ws.cell(row=grand_total_row - 1, column=1, value="═" * 50)
            ws.cell(row=grand_total_row - 1, column=1).font = Font(name='Courier New', size=10, color='FF0000')
            
            # Grand total with order count
            ws.cell(row=grand_total_row, column=2, value=f"🏆 GRAND TOTAL ({order_count} orders):")
            ws.cell(row=grand_total_row, column=2).font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
            ws.cell(row=grand_total_row, column=2).alignment = Alignment(horizontal='right')
            ws.cell(row=grand_total_row, column=2).fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            
            ws.cell(row=grand_total_row, column=4, value=grand_total)
            ws.cell(row=grand_total_row, column=4).font = grand_total_font
            ws.cell(row=grand_total_row, column=4).alignment = Alignment(horizontal='right')
            ws.cell(row=grand_total_row, column=4).fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            ws.cell(row=grand_total_row, column=4).border = thick_border
            
            # Summary section
            summary_row = grand_total_row + 2
            ws.cell(row=summary_row, column=1, value=f"📊 Invoice Summary for {customer}:")
            ws.cell(row=summary_row, column=1).font = Font(name='Arial', size=12, bold=True, color='4472C4')
            
            ws.cell(row=summary_row + 1, column=1, value=f"• Total Orders: {order_count}")
            ws.cell(row=summary_row + 2, column=1, value=f"• Latest Order: {current_date} at {current_time}")
            if existing_files:
                file_creation_date = datetime.fromtimestamp(os.path.getctime(existing_file)).strftime('%d/%m/%Y')
                ws.cell(row=summary_row + 3, column=1, value=f"• First Order: {file_creation_date}")
            
            # Format currency
            for row_num in range(new_start_row + 1, row):
                ws.cell(row=row_num, column=3).number_format = '"$"#,##0.00'
                ws.cell(row=row_num, column=4).number_format = '"$"#,##0.00'
            ws.cell(row=total_row, column=4).number_format = '"$"#,##0.00'
            ws.cell(row=grand_total_row, column=4).number_format = '"$"#,##0.00'
            
            # Column widths
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 25
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 15
            
            # Save file
            full_path = os.path.abspath(filename)
            print(f"DEBUG: Saving to: {full_path}")
            
            wb.save(filename)
            
            print(f"DEBUG: File saved successfully")
            
            # Verify file exists
            if os.path.exists(filename):
                size = os.path.getsize(filename)
                print(f"DEBUG: File exists, size: {size} bytes")
                
                if existing_files:
                    success_msg = (f"Order ADDED to existing invoice!\n\n"
                                 f"File: {filename}\n"
                                 f"Location: {full_path}\n"
                                 f"Customer: {customer}\n"
                                 f"This Order: ${order_total:.2f}\n"
                                 f"📊 GRAND TOTAL: ${grand_total:.2f}\n"
                                 f"📅 Total Orders: {order_count}")
                else:
                    success_msg = (f"NEW invoice created!\n\n"
                                 f"File: {filename}\n"
                                 f"Location: {full_path}\n"
                                 f"Customer: {customer}\n"
                                 f"📊 Total: ${grand_total:.2f}")
                
                messagebox.showinfo("Success!", success_msg)
                self.status_label.config(text=f"Invoice saved: {filename}", foreground='green')
                
                # Reset form for next order
                self.reset_products_only()
            else:
                print("DEBUG: File was not created!")
                messagebox.showerror("Error", "File was not created!")
                
        except Exception as e:
            error_msg = f"Error creating invoice: {str(e)}"
            print(f"DEBUG ERROR: {error_msg}")
            messagebox.showerror("Error", error_msg)
            
    def reset_products_only(self):
        """Reset only product fields, keep customer info for multiple orders"""
        self.product_entry.delete(0, tk.END)
        self.qty_entry.delete(0, tk.END)
        self.price_entry.delete(0, tk.END)
        self.product_listbox.delete(0, tk.END)
        self.products.clear()
        self.update_total()
        self.status_label.config(text="Ready for next order...", foreground='blue')
            
    def reset_form(self):
        """Reset the form for new customer"""
        self.shop_entry.delete(0, tk.END)
        self.customer_entry.delete(0, tk.END)
        self.area_entry.delete(0, tk.END)
        self.product_entry.delete(0, tk.END)
        self.qty_entry.delete(0, tk.END)
        self.price_entry.delete(0, tk.END)
        self.product_listbox.delete(0, tk.END)
        self.products.clear()
        self.update_total()
        self.status_label.config(text="Ready for new invoice...", foreground='blue')


def main():
    root = tk.Tk()
    app = SimpleInvoiceApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()